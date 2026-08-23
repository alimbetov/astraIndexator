from __future__ import annotations

import re
from dataclasses import replace
from pathlib import Path
from typing import Iterable

import pdfplumber
from docx import Document
from PIL import Image

from astra_indexator.acquisition import AcquiredSource

from .base import ParseContext, ParserError
from .model import (
    DocumentElement,
    ElementType,
    OcrCandidate,
    ParsedDocument,
    ParseQuality,
    ParserIdentity,
    QualityStatus,
    SourceGeometry,
    deterministic_element_id,
)


def _eid(source: AcquiredSource, ctx: ParseContext, locator: str, kind: ElementType) -> str:
    return deterministic_element_id(
        document_id=ctx.document_id,
        document_version=ctx.document_version,
        source_sha256=source.sha256,
        detected_format=source.detected_format,
        locator=locator,
        element_type=kind,
    )


def _quality(elements: Iterable[DocumentElement], candidates: list[OcrCandidate], warnings: list[str], page_modes=()) -> ParseQuality:
    chars = sum(len(e.text or "") for e in elements)
    if candidates:
        status = QualityStatus.OCR_REQUIRED
    elif chars == 0:
        status = QualityStatus.LOW_SIGNAL
    else:
        status = QualityStatus.GOOD
    return ParseQuality(status, chars, len(candidates), tuple(warnings), tuple(page_modes))


def _doc(source: AcquiredSource, ctx: ParseContext, name: str, elements: list[DocumentElement], candidates: list[OcrCandidate], warnings: list[str], page_modes=()) -> ParsedDocument:
    if len(elements) > ctx.limits.max_elements:
        raise ParserError("PARSER_ELEMENT_LIMIT", "parser element limit exceeded")
    chars = sum(len(e.text or "") for e in elements)
    if chars > ctx.limits.max_extracted_chars:
        raise ParserError("PARSER_TEXT_LIMIT", "parser extracted character limit exceeded")
    return ParsedDocument(
        schema_version="astra-indexator-document-v1",
        document_id=ctx.document_id,
        document_version=ctx.document_version,
        source_sha256=source.sha256,
        detected_format=source.detected_format,
        parser=ParserIdentity(name=name, version="m4-v1", profile=ctx.parser_profile),
        elements=tuple(elements),
        ocr_candidates=tuple(candidates),
        quality=_quality(elements, candidates, warnings, page_modes),
    )


class TextDocumentHandler:
    format_name = "TXT"

    def parse(self, source: AcquiredSource, context: ParseContext) -> ParsedDocument:
        text = source.local_path.read_text(encoding="utf-8")
        blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
        elements = [
            DocumentElement(
                element_id=_eid(source, context, f"paragraph:{i}", ElementType.PARAGRAPH),
                type=ElementType.PARAGRAPH,
                order_index=i,
                text=block,
                source_locator={"paragraphIndex": i},
            )
            for i, block in enumerate(blocks)
        ]
        return _doc(source, context, "text-native", elements, [], [])


class MarkdownDocumentHandler:
    format_name = "MARKDOWN"

    def parse(self, source: AcquiredSource, context: ParseContext) -> ParsedDocument:
        lines = source.local_path.read_text(encoding="utf-8").splitlines()
        elements: list[DocumentElement] = []
        section_path: list[str] = []
        in_code = False
        code: list[str] = []
        paragraph: list[str] = []

        def emit(kind: ElementType, text: str, locator: str, *, level: int | None = None) -> None:
            idx = len(elements)
            elements.append(DocumentElement(
                element_id=_eid(source, context, locator, kind), type=kind, order_index=idx, text=text,
                level=level, section_path=tuple(section_path), source_locator={"lineLocator": locator}
            ))

        def flush_paragraph(line_no: int) -> None:
            if paragraph:
                emit(ElementType.PARAGRAPH, "\n".join(paragraph).strip(), f"paragraph-ending-line:{line_no}")
                paragraph.clear()

        for line_no, line in enumerate(lines, start=1):
            if line.strip().startswith("```"):
                if in_code:
                    emit(ElementType.CODE_BLOCK, "\n".join(code), f"code-ending-line:{line_no}")
                    code.clear()
                    in_code = False
                else:
                    flush_paragraph(line_no)
                    in_code = True
                continue
            if in_code:
                code.append(line)
                continue
            heading = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading:
                flush_paragraph(line_no)
                level = len(heading.group(1)); title = heading.group(2).strip()
                section_path[:] = section_path[: level - 1]
                section_path.append(title)
                emit(ElementType.HEADING, title, f"heading-line:{line_no}", level=level)
                continue
            item = re.match(r"^\s*(?:[-*+] |\d+[.)]\s+)(.+)$", line)
            if item:
                flush_paragraph(line_no)
                emit(ElementType.LIST_ITEM, item.group(1).strip(), f"list-line:{line_no}")
                continue
            if not line.strip():
                flush_paragraph(line_no)
            else:
                paragraph.append(line)
        flush_paragraph(len(lines) + 1)
        if in_code and code:
            emit(ElementType.CODE_BLOCK, "\n".join(code), "code-eof")
        return _doc(source, context, "markdown-native", elements, [], [])


class DocxDocumentHandler:
    format_name = "DOCX"

    def parse(self, source: AcquiredSource, context: ParseContext) -> ParsedDocument:
        docx = Document(str(source.local_path))
        elements: list[DocumentElement] = []
        section_path: list[str] = []
        table_count = 0
        for p_idx, paragraph in enumerate(docx.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            match = re.match(r"Heading\s+(\d+)", style, re.IGNORECASE)
            kind = ElementType.HEADING if match else ElementType.PARAGRAPH
            level = int(match.group(1)) if match else None
            if kind is ElementType.HEADING:
                section_path[:] = section_path[: (level or 1) - 1]
                section_path.append(text)
            elements.append(DocumentElement(
                element_id=_eid(source, context, f"paragraph:{p_idx}", kind),
                type=kind, order_index=len(elements), text=text, level=level,
                section_path=tuple(section_path), source_locator={"paragraphIndex": p_idx, "style": style},
            ))
        for t_idx, table in enumerate(docx.tables):
            table_count += 1
            if table_count > context.limits.max_tables:
                raise ParserError("PARSER_TABLE_LIMIT", "table limit exceeded")
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            elements.append(DocumentElement(
                element_id=_eid(source, context, f"table:{t_idx}", ElementType.TABLE),
                type=ElementType.TABLE, order_index=len(elements), text=None,
                section_path=tuple(section_path), source_locator={"tableIndex": t_idx},
                metadata={"rows": rows, "rowCount": len(rows), "columnCount": max((len(r) for r in rows), default=0)},
            ))
        return _doc(source, context, "docx-native", elements, [], [])


class ImageDocumentHandler:
    def __init__(self, format_name: str):
        self.format_name = format_name

    def parse(self, source: AcquiredSource, context: ParseContext) -> ParsedDocument:
        with Image.open(source.local_path) as image:
            frames = getattr(image, "n_frames", 1)
            elements: list[DocumentElement] = []
            candidates: list[OcrCandidate] = []
            for frame in range(frames):
                locator = f"frame:{frame + 1}"
                eid = _eid(source, context, locator, ElementType.IMAGE)
                geom = SourceGeometry(page_number=frame + 1, page_width=float(image.width), page_height=float(image.height), coordinate_space="pixels")
                elements.append(DocumentElement(eid, ElementType.IMAGE, len(elements), geometry=geom, source_locator={"frame": frame + 1}, role="SCANNED_PAGE"))
                candidates.append(OcrCandidate(f"ocr:{eid}", "PAGE", frame + 1, "standalone_image", eid, geom))
            return _doc(source, context, f"{self.format_name.lower()}-image-shell", elements, candidates, [], ["SCANNED_IMAGE"] * frames)


class PdfDocumentHandler:
    format_name = "PDF"

    def parse(self, source: AcquiredSource, context: ParseContext) -> ParsedDocument:
        elements: list[DocumentElement] = []
        candidates: list[OcrCandidate] = []
        warnings: list[str] = []
        page_modes: list[str] = []
        with pdfplumber.open(str(source.local_path)) as pdf:
            if len(pdf.pages) > context.limits.max_pages:
                raise ParserError("PARSER_PAGE_LIMIT", "PDF page limit exceeded")
            for page_number, page in enumerate(pdf.pages, start=1):
                words = page.extract_words(use_text_flow=False, keep_blank_chars=False) or []
                text_chars = sum(len(str(w.get("text", ""))) for w in words)
                images = page.images or []
                page_area = max(float(page.width) * float(page.height), 1.0)
                image_area = sum(max(0.0, float(i.get("x1", 0)) - float(i.get("x0", 0))) * max(0.0, float(i.get("bottom", 0)) - float(i.get("top", 0))) for i in images)
                image_ratio = min(image_area / page_area, 1.0)
                if text_chars >= context.quality_profile.min_good_chars_per_page:
                    mode = "MIXED" if images else "NATIVE_TEXT"
                elif images and image_ratio >= context.quality_profile.scanned_image_area_ratio:
                    mode = "SCANNED_IMAGE"
                elif text_chars <= context.quality_profile.low_signal_chars_per_page:
                    mode = "LOW_SIGNAL" if words or images else "EMPTY"
                else:
                    mode = "LOW_SIGNAL"
                page_modes.append(mode)

                # Deterministic v1 order: words sorted top-to-bottom then left-to-right.
                # Column-aware refinements remain profile-versioned; no native object order leaks downstream.
                ordered = sorted(words, key=lambda w: (round(float(w.get("top", 0)), 1), round(float(w.get("x0", 0)), 1)))
                lines: list[list[dict]] = []
                for word in ordered:
                    top = float(word.get("top", 0))
                    if not lines or abs(top - float(lines[-1][0].get("top", 0))) > 4.0:
                        lines.append([word])
                    else:
                        lines[-1].append(word)
                for line_idx, line in enumerate(lines):
                    line = sorted(line, key=lambda w: float(w.get("x0", 0)))
                    text = " ".join(str(w.get("text", "")) for w in line).strip()
                    if not text:
                        continue
                    x0 = min(float(w.get("x0", 0)) for w in line); x1 = max(float(w.get("x1", 0)) for w in line)
                    top = min(float(w.get("top", 0)) for w in line); bottom = max(float(w.get("bottom", 0)) for w in line)
                    locator = f"page:{page_number}:line:{line_idx}"
                    geom = SourceGeometry(page_number, x0, top, x1, bottom, float(page.width), float(page.height), "pdf-points-top-left")
                    elements.append(DocumentElement(
                        _eid(source, context, locator, ElementType.PARAGRAPH), ElementType.PARAGRAPH, len(elements), text=text,
                        geometry=geom, source_locator={"pageNumber": page_number, "lineIndex": line_idx},
                    ))

                for image_idx, image in enumerate(images):
                    locator = f"page:{page_number}:image:{image_idx}"
                    geom = SourceGeometry(page_number, float(image.get("x0", 0)), float(image.get("top", 0)), float(image.get("x1", 0)), float(image.get("bottom", 0)), float(page.width), float(page.height), "pdf-points-top-left")
                    eid = _eid(source, context, locator, ElementType.IMAGE)
                    elements.append(DocumentElement(eid, ElementType.IMAGE, len(elements), geometry=geom, source_locator={"pageNumber": page_number, "imageIndex": image_idx}, role="UNKNOWN"))
                    candidates.append(OcrCandidate(f"ocr:{eid}", "EMBEDDED_IMAGE", page_number, "embedded_image", eid, geom))

                if mode in {"SCANNED_IMAGE", "LOW_SIGNAL"}:
                    locator = f"page:{page_number}:ocr"
                    candidates.append(OcrCandidate(f"ocr-page:{page_number}", "PAGE", page_number, mode.lower(), metadata={"pageMode": mode}))
                elements.append(DocumentElement(
                    _eid(source, context, f"page:{page_number}:break", ElementType.PAGE_BREAK), ElementType.PAGE_BREAK,
                    len(elements), source_locator={"pageNumber": page_number}
                ))
        return _doc(source, context, "pdfplumber-layout", elements, candidates, warnings, page_modes)
